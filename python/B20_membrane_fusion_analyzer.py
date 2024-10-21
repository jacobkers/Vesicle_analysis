import numpy as np
import csv
from pathlib import Path
import matplotlib.pyplot as plt
from openpyxl import load_workbook
from scipy.optimize import curve_fit

# exponential function with background
def exponential_model(t, A, k, B, t0):
    return A * np.exp(-k * (t - t0)) + B

class Event:
    def __init__(self):
        self.index = 0
        self.x0 = 0
        self.y0 = 0
        self.t0 = 0
        self.type = 0

# Example usage
if 0: 
    label='2_TIRF_488_001_PCPG_Chol_small_short'
    moviepath=Path('M:/tnw/bn/cd\Shared/Jacob/TESTdata_in/2023_Rafa/2024_10_02 membrane fusion')
    savepath=Path('M:/tnw/bn/cd\Shared/Jacob/TESTdata_out/2023_Rafa/2024_10_02 membrane fusion')
    movie_filename = '2_TIRF_488_001_PCPG_Chol_small_short.tif'
    filename ='STD_2_TIRF_488_001_PCPG_Chol_small_short.tif'
    xls_classification = savepath / "Events_classification_jacob_short.xlsx"  
if 1:
    label='2_TIRF_488_001_PCPG_Chol_long'
    moviepath=Path('M:/tnw/bn/cd\Shared/Jacob/TESTdata_in/2023_Rafa/2024_10_02 membrane fusion')
    savepath=Path('M:/tnw/bn/cd\Shared/Jacob/TESTdata_out/2023_Rafa/2024_10_02 membrane fusion')
    movie_filename = '2_TIRF_488_001_PCPG_Chol.tif'
    filename ='STD_2_TIRF_488_001_PCPG_Chol.tif'
    xls_classification  = savepath / "Events_classification_jacob.xlsx"  
if 0:
    moviepath= Path.cwd()
    movie_filename = '40 uM LUVsWITHCerC6_RealTime_Series002_t000_crp-1_red.tif'
    filename = 'MAX_40 uM LUVsWITHCerC6_RealTime_Series002_t000_crp-1.tif (red).tif'

image_path =  moviepath/ filename
movie_path =  moviepath/ movie_filename


#load traces
trace_data_name=label +"_traces" +  str(".csv")
csv_traces=savepath /  trace_data_name
csv_path_in = Path(csv_traces)
print(csv_traces.stem)
pre_trace_data = np.loadtxt(csv_traces, delimiter=';')

#load classification file of events (contains t0,x,y,type)
xls_source = savepath / xls_classification 
wb = load_workbook(filename = xls_classification)
sheet_events = wb['events']
ColNames = {}
Current  = 0
for COL in sheet_events.iter_cols(1, sheet_events.max_column):
    ColNames[COL[0].value] = Current
    Current += 1
event_list=[]
for row_cells in sheet_events.iter_rows(min_row=2):
    event=Event()
    event.type=((row_cells[ColNames['type']].value))
    event.x0=((row_cells[ColNames['x']].value))
    event.y0=((row_cells[ColNames['y']].value))
    event.t0=((row_cells[ColNames['t0']].value))
    event_list.append(event)

#we need to add:
# dark treshold
# end-of-trace residu
# dock time
# event starts (maxima)
# release times (minima)
# end-of-event (=eot or back-to-dark)

# plot traces and start_time
xls_source = savepath / xls_source 
frs,N_events=np.shape(pre_trace_data)

fig, ax=plt.subplots(2,2)
for event,pre_trace in zip(event_list,pre_trace_data.T):
    dif_trace=np.diff(pre_trace)
    t0=int(event.t0)
    tp=event.type   
    start=np.max([t0-20, 0])
    stop=np.min([t0+100, frs])
    trace_cut=pre_trace[start:stop]
    dif_trace_cut=dif_trace[start:stop]
    t1=np.argmin(dif_trace)
    
    
    #show:
    if tp == "hemifusion":
        ax[0,0].plot(pre_trace, '-', markersize=2)
        ax[0,0].plot(t0,pre_trace[t0], 'ro-', markersize=4)
        ax[0,0].set_title('traces')
        ax[0,1].plot(dif_trace, '-', markersize=2)
        ax[0,1].plot(t0, dif_trace[t0], 'ro-', markersize=4)
        ax[0,1].plot(t1, dif_trace[t1], 'go-', markersize=4)
        ax[0,1].set_title('derivative')
        ax[1,0].plot(trace_cut, '-')
        ax[1,0].set_title('aligned')
        ax[1,1].plot(trace_cut/np.max(trace_cut), '-')
        ax[1,1].set_title('normalized')
        fig.show()
        dum=1
plt.close('all')


    







