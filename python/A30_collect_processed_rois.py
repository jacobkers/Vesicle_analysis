"""
21-2-2024
Work guv imagery
@author: jkerssemakers
"""
import csv
import numpy as np
from pathlib import Path
from A00_init import get_exps
import matplotlib.pyplot as plt
from openpyxl import load_workbook
from scipy.interpolate import make_interp_spline

""" 
experiment indices (int = laptop, add: 0.1 for office local, 0.2 to run on CD:K:):
0: .nd testfiles
1: .lif testfiles
2: .tif testfiles (only laptop)
3: .tif test (office-PC) less_challenging ones

"""



class GUV:
    """
    properties of a single guv movie
    """
    def __init__(self):
            #priority	identifier	comment	format	nZ	dT [s]	path in_order_of_appearance
            self.exp_id=0
            self.label='any_label'
            self.comment=[]
            self.use_it=[]
            self.crop_it=[]
            self.dt=1       

def get_data_selections(run_id):
    #use a local copy:
    wb = load_workbook(filename =  "data_overview.xlsx")
    sheet_files = wb['guvs']
    #Create a dictionary of column names
    
    Header = {}
    Current  = 0
    for COL in sheet_files.iter_cols(1, sheet_files.max_column):
        Header[COL[0].value] = Current
        Current += 1
    Guv_list=[]

    mx_row=sheet_files.max_row
    for row_cells in sheet_files.iter_rows(min_row=2, max_row=mx_row):
        Guv = GUV()
        Guv.exp_id=(row_cells[Header["exp_id"]].value)
        Guv.movie_id=(row_cells[Header["movie_id"]].value)
        Guv.label=(row_cells[Header["guv_label"]].value) 
        Guv.use_it=(row_cells[Header["use"]].value) 
        Guv.crop_it=(row_cells[Header["crop"]].value)
        Guv.dt=(row_cells[Header["dt(s)"]].value)
        if Guv.exp_id == run_id:
            Guv_list.append(Guv)
    return Guv_list
 

expi = 5.2   #2: flexibles; 3:less_challenging ones 4: single-image tiffs
movie_to_use_list=[462, 463,465,466]
movie_simbol_list=[str('r'), str('k'), str('m'),str('b')]

initval = get_exps(expi)

if initval.suffix =='.tif'and initval.sequence=='single_frame':
    #collect single-frame data points and re-save all in one file
    load_dirname=initval.mainpath_out + initval.subdir +str("/A20b_processed/")
    csv_path_in = Path(load_dirname)
    name=[]
    X=[]
    Y=[]
    data=[]
    for csv_source in csv_path_in.glob("**/*.csv"):  # find all relevant csv files in inpath   
        print(csv_source.stem)
        with open(csv_source) as f:
            reader = csv.DictReader(f, delimiter=";")
            for row in reader:
                name.append(csv_source.stem)
                data.append(row)
    #set up csv for tracking data:
    csv_path_out = Path(initval.mainpath_out + initval.subdir +str("/A30_processed/"))
    if not csv_path_out.is_dir():
            csv_path_out.mkdir()
            
    csv_target=load_dirname  + "collected_data" + str(".csv")
    with open(csv_target, "w",newline='') as csv_f:  # will overwrite existing
        # create the csv writer
        writer = csv.writer(csv_f, delimiter=";")
        #f = open("test.csv", "a")
        writer.writerow(row.keys())
        for data_row in data:  
            
            # create the csv writer
            writer = csv.writer(csv_f, delimiter=";")
            #f = open("test.csv", "a")
            writer.writerow(data_row.values()) 
        dum=1
    

if initval.suffix =='.tif'and initval.sequence=='time_trace':
    fig, axs=plt.subplots(2,2)
     #collect files per trace:
    Guv_list=get_data_selections(round(expi))
    
    load_dirname_A20a=initval.mainpath_out + initval.subdir +str("/A20a_tracked/")
    load_dirname_A20b=initval.mainpath_out + initval.subdir +str("/A20b_processed/")
    csv_path_in_A20a = Path(load_dirname_A20a)
    csv_path_in_A20b = Path(load_dirname_A20b)
    guv_labels=[]
    #select movie and guvs:
    for this_guv in Guv_list:
        if this_guv.use_it ==1 and this_guv.movie_id in movie_to_use_list:
            movie_use_index = movie_to_use_list.index(this_guv.movie_id)
            
            guv_labels.append(this_guv.label)
            data=[]
            csv_source=Path(load_dirname_A20b + this_guv.label +"_all_data.csv")
            with open(csv_source) as g:
                reader = csv.DictReader(g, delimiter=";")
                for row in reader:
                    data.append(row)
            #collect and plot
  
            plot_ax=[]
            area=[]
            roundness=[]
            major_minor=[]
            c0_edge_mx=[]
            c1_edge_mx=[]
            c0_edge_sm=[]
            c1_edge_sm=[]
            c1_std_sum_rel=[]
            ratio1=[]
            for fri, row in enumerate(data):
                c1_signal_check=(float(row['c1_edge_mx']))
                okay_point=float(row['roundness'])>0 and float(row['R_minor'])>0 and c1_signal_check>400 and fri< this_guv.crop_it
                if okay_point:
                    #fetch work parameters:
                    axis_major=float(row['R_major'])
                    axis_minor=float(row['R_minor'])
                    c1_mean_peaks=(float(row['c1_edge_mx']))
                    #c0_mean_peaks=(float(row['c0_edge_mx']))
                    #build plots in calibrated units
                    #choose axis and build:
                    if 1:
                        plot_ax.append(this_guv.dt*fri)
                        axlabel="time (s)"
                    if 0:
                        plot_ax.append(initval.counts2perc*c1_mean_peaks)
                        axlabel="content (%)"
                    #build parameters
                    area.append((initval.pix2mu)**2*float(row['area']))
                    roundness.append(float(row['roundness']))
                    major_minor.append(axis_major/axis_minor)
                    c1_edge_mx.append(initval.counts2perc*c1_mean_peaks)
         
            simbol_color=movie_simbol_list[movie_use_index]
            sz=2


            axs[0,0].plot(plot_ax, area,'o-', markersize=sz, color=simbol_color)
            axs[0,0].set_ylabel("area, mu^2")
            axs[0,0].set_xlabel(axlabel)
            axs[0,0].set_title("area")  

            axs[0,1].plot(plot_ax,roundness, 'o-', markersize=sz, color=simbol_color)
            axs[0,1].set_xlabel(axlabel)
            axs[0,1].set_ylabel("roundness")
            axs[0,1].set_title("roundness")  

            axs[1,0].plot(plot_ax,c1_edge_mx, 'o-', markersize=sz, color=simbol_color)
            axs[1,0].set_xlabel(axlabel)
            axs[1,0].set_ylabel("content (%)")
            axs[1,0].set_title("edge content") 
            axs[1,0].legend(guv_labels,loc='best', fontsize='xx-small') 

            axs[1,1].plot(plot_ax, major_minor, 'o-', markersize=sz, color=simbol_color)
            axs[1,1].set_xlabel(axlabel)
            axs[1,1].set_ylabel("ratio, a.u.")
            axs[1,1].set_title("major/minor")
         
    fig.show()
    dum=1


